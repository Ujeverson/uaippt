"""Gera assets/template.pptx a partir do template original do SENAI (SlideSENAI.pptx).

Uso (manutenção da skill — só é necessário se o template original mudar):
    python build_template.py <SlideSENAI.pptx> <saida/template.pptx> [<saida/fatesg_rodape.png>]

Ajustes aplicados sobre o original:
  1. Placeholders de título/conteúdo deslocados para baixo do banner de circuitos
     (o banner ocupa y=0.00"–1.00"; no original o título começava em y=0.40").
  2. Imagem de rodapé da Faculdade SENAI Fatesg embutida no layout de capa
     ("Slide de Título"), para que toda capa a receba automaticamente.
  3. Cores do tema trocadas para a paleta da marca SENAI (títulos em azul SENAI).
  4. Slides de exemplo removidos (o template sai com 0 slides).
"""
import copy
import sys

from lxml import etree
from pptx import Presentation
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.util import Inches

NS = {
    "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
    "p": "http://schemas.openxmlformats.org/presentationml/2006/main",
    "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
}

# Paleta da marca (amostrada da logo e do banner do template)
THEME_COLORS = {
    "dk2": "0E4194",      # azul SENAI (títulos)
    "lt2": "E8EEF7",      # azul muito claro (fundos de caixa)
    "accent1": "0E4194",  # azul SENAI
    "accent2": "E8490F",  # laranja SENAI
    "accent3": "074982",  # azul petróleo (linhas do banner)
    "accent4": "4FB3D9",  # ciano (setas do banner)
    "accent5": "575656",  # cinza da logo
    "accent6": "2E8B57",  # verde (status/positivo)
}

# Geometria nova, em polegadas: (x, y, w, h). Chave = (nome do layout, idx do placeholder)
TITLE_STD = (0.92, 1.05, 11.50, 0.90)
GEOMETRY = {
    "MASTER": {0: TITLE_STD, 1: (0.92, 2.05, 11.50, 4.75)},
    "Slide de Título": {0: (1.67, 1.45, 10.00, 2.45), 1: (1.67, 4.15, 10.00, 1.50)},
    "Título e Conteúdo": {0: TITLE_STD, 1: (0.92, 2.05, 11.50, 4.75)},
    "Duas Partes de Conteúdo": {0: TITLE_STD, 1: (0.92, 2.05, 5.67, 4.75), 2: (6.75, 2.05, 5.67, 4.75)},
    "Comparação": {
        0: TITLE_STD,
        1: (0.92, 2.05, 5.64, 0.60), 2: (0.92, 2.70, 5.64, 4.10),
        3: (6.75, 2.05, 5.67, 0.60), 4: (6.75, 2.70, 5.67, 4.10),
    },
    "Somente Título": {0: TITLE_STD},
    "Conteúdo com Legenda": {0: (0.92, 1.15, 4.30, 1.20), 1: (5.67, 1.15, 6.75, 5.65), 2: (0.92, 2.45, 4.30, 4.35)},
    "Imagem com Legenda": {0: (0.92, 1.15, 4.30, 1.20), 1: (5.67, 1.15, 6.75, 5.65), 2: (0.92, 2.45, 4.30, 4.35)},
    "Título e Texto Vertical": {0: TITLE_STD, 1: (0.92, 2.05, 11.50, 4.75)},
    "Texto e Título Vertical": {0: (9.54, 1.10, 2.88, 5.70), 1: (0.92, 1.10, 8.46, 5.70)},
}


def set_geom(shape, geom):
    x, y, w, h = (Inches(v) for v in geom)
    shape.left, shape.top, shape.width, shape.height = x, y, w, h


def apply_geometry(container, table):
    for ph in container.placeholders:
        g = table.get(ph.placeholder_format.idx)
        if g:
            set_geom(ph, g)


def main(src, dst, logo_out=None):
    prs = Presentation(src)
    master = prs.slide_masters[0]

    # 1) Geometria
    apply_geometry(master, GEOMETRY["MASTER"])
    for layout in prs.slide_layouts:
        if layout.name in GEOMETRY:
            apply_geometry(layout, GEOMETRY[layout.name])

    # 2) Rodapé Fatesg (imagem do slide 1 original) -> layout de capa
    cover_slide = prs.slides[0]
    pic = next(s for s in cover_slide.shapes if s.shape_type == 13)
    rid_old = pic._element.blipFill.find(f"{{{NS['a']}}}blip").get(f"{{{NS['r']}}}embed")
    image_part = cover_slide.part.related_part(rid_old)
    if logo_out:
        with open(logo_out, "wb") as f:
            f.write(image_part.blob)
    cover_layout = next(l for l in prs.slide_layouts if l.name == "Slide de Título")
    rid_new = cover_layout.part.relate_to(image_part, RT.IMAGE)
    new_pic = copy.deepcopy(pic._element)
    new_pic.blipFill.find(f"{{{NS['a']}}}blip").set(f"{{{NS['r']}}}embed", rid_new)
    new_pic.nvPicPr.cNvPr.set("name", "Rodape Fatesg")
    new_pic.nvPicPr.cNvPr.set("id", "99")
    cover_layout.shapes._spTree.append(new_pic)

    # 3) Cores do tema + títulos em azul SENAI (tx2 = dk2)
    theme_part = master.part.part_related_by(RT.THEME)
    theme = etree.fromstring(theme_part.blob)
    scheme = theme.find(".//a:clrScheme", NS)
    scheme.set("name", "SENAI")
    for tag, hexval in THEME_COLORS.items():
        el = scheme.find(f"a:{tag}", NS)
        for child in list(el):
            el.remove(child)
        etree.SubElement(el, f"{{{NS['a']}}}srgbClr", val=hexval)
    theme_part._blob = etree.tostring(theme, xml_declaration=True, encoding="UTF-8", standalone=True)

    title_fill = master._element.find(".//p:txStyles/p:titleStyle/a:lvl1pPr/a:defRPr/a:solidFill", NS)
    for child in list(title_fill):
        title_fill.remove(child)
    etree.SubElement(title_fill, f"{{{NS['a']}}}schemeClr", val="tx2")

    # 4) Remove slides de exemplo
    while len(prs.slides) > 0:
        rId = prs.slides._sldIdLst[0].rId
        prs.part.drop_rel(rId)
        del prs.slides._sldIdLst[0]

    prs.save(dst)
    print(f"OK -> {dst}")


if __name__ == "__main__":
    if len(sys.argv) < 3:
        sys.exit(__doc__)
    main(*sys.argv[1:4])
