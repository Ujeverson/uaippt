"""uaippt_deck — gera apresentações .pptx no padrão visual SENAI / Faculdade SENAI Fatesg.
Suporta Motion Design com Manim e Blender (bpy) e integração com AI Coding Agents.

Padrão tipográfico (baseado em slideExemploComConteudo.pdf):
    Capa .............. título Poppins 36 · nome do(a) professor(a) Poppins 40
    Slides 2 em diante  título Montserrat 36 · conteúdo Calibri 24
    Disposição ........ título sob o banner, texto à esquerda, imagem à direita
                        (crédito/legenda 12 pt abaixo da imagem), @perfil no canto inferior esquerdo.

Uso como biblioteca:
    import sys; sys.path.insert(0, r"<pasta-da-skill>/scripts")
    from uaippt_deck import UaipptDeck

    deck = UaipptDeck(social_handle="@prof.ujeverson")
    deck.cover("Lógica de Programação", "Aula 1 — Algoritmos", course="Análise e Des. de Sistemas",
               professor="Prof. Msc. Ujeverson Tavares")
    deck.bullets("Objetivos da aula", ["Entender o que é um algoritmo", ("Exemplos do dia a dia", 1)])
    deck.save("aula1.pptx")

Uso via linha de comando (roteiro em JSON):
    python uaippt_deck.py roteiro.json saida.pptx

Dependências: python-pptx (obrigatória), pygments (opcional, colore código).
As fontes Poppins, Montserrat e Calibri devem estar instaladas no computador que abrir o .pptx.
"""
from __future__ import annotations

import json
import math
import re
import sys
import warnings
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt
from lxml import etree

SKILL_DIR = Path(__file__).resolve().parent.parent
TEMPLATE = SKILL_DIR / "assets" / "template.pptx"
SOCIAL_ICON = SKILL_DIR / "assets" / "instagram_icon.png"

# ---------------------------------------------------------------- paleta
AZUL = RGBColor(0x0E, 0x41, 0x94)        # azul SENAI
LARANJA = RGBColor(0xE8, 0x49, 0x0F)     # laranja SENAI
PETROLEO = RGBColor(0x07, 0x49, 0x82)    # linhas do banner
CIANO = RGBColor(0x4F, 0xB3, 0xD9)       # setas do banner
CINZA = RGBColor(0x57, 0x56, 0x56)       # cinza da logo
TEXTO = RGBColor(0x1A, 0x1A, 0x1A)       # texto corrido
AZUL_CLARO = RGBColor(0xE8, 0xEE, 0xF7)  # fundo de caixas/linhas alternadas
BRANCO = RGBColor(0xFF, 0xFF, 0xFF)
CODE_BG = RGBColor(0x1E, 0x1E, 0x2E)
CODE_FG = RGBColor(0xE6, 0xE6, 0xE6)

# ---------------------------------------------------------------- tipografia (padrão do slide de exemplo)
FONT_COVER = "Poppins"
FONT_TITLE = "Montserrat"
FONT_BODY = "Calibri"
COVER_TITLE_PT = 36
COVER_PROF_PT = 40
TITLE_PT = 36
BODY_PT = 24
CAPTION_PT = 12

DEFAULT_PROFESSOR = ""
DEFAULT_COURSE = "Faculdade SENAI Fatesg"

# ---------------------------------------------------------------- layouts (índices do template)
L_CAPA, L_CONTEUDO, L_SECAO, L_DUAS, L_COMPARACAO = 0, 1, 2, 3, 4
L_SO_TITULO, L_BRANCO, L_CONTEUDO_LEGENDA, L_IMAGEM_LEGENDA = 5, 6, 7, 8

# Geometria (polegadas) medida no slide de exemplo (960 × 540 pt)
LEFT = 0.66                                   # caixa de texto; o texto começa em 0,76"
TITLE_BOX = dict(x=LEFT, y=0.98, w=12.0, h=0.78)
AREA = dict(x=LEFT, y=1.85, w=12.0, h=4.95)    # área de conteúdo (termina em y = 6,80")
TEXT_W_WITH_IMG = 7.25                         # largura do texto quando há imagem à direita
IMG_BOX = dict(x=8.25, y=1.80, w=4.80, h=4.45) # coluna da imagem à direita


# ---------------------------------------------------------------- utilidades de texto
def _norm_items(items):
    """Aceita 'texto', ('texto', nivel) ou ['texto', nivel]; devolve [(texto, nivel)]."""
    out = []
    for it in items or []:
        if isinstance(it, (list, tuple)):
            out.append((str(it[0]), int(it[1]) if len(it) > 1 else 0))
        else:
            out.append((str(it), 0))
    return out


def _plain(text):
    return text.replace("**", "")


def _text_height(items, width_in, size=BODY_PT, bulleted=True):
    """Altura estimada (pol.) do texto em Calibri; usada só para avisar de estouro."""
    total = 0.12
    for text, level in items:
        s = size
        indent = 0.35 * (level + 1) if (bulleted or level) else 0.0
        usable = width_in - 0.2 - indent
        cpl = max(8, int(usable * 72 / (s * 0.49)))
        lines = max(1, math.ceil(len(_plain(text)) / cpl))
        total += lines * s * 1.2 / 72 + 6 / 72
    return total


def _warn_overflow(label, items, w, h, size=BODY_PT, bulleted=True):
    need = _text_height(items, w, size, bulleted)
    if need > h + 0.05:
        warnings.warn(f"[uaippt] '{label}': texto (~{need:.1f}\") excede a área ({h:.1f}\"). "
                      f"Divida o slide — o padrão exige Calibri {size}.", stacklevel=3)


def _set_font(run, name, size, bold=None, italic=None, color=None):
    run.font.name = name
    run.font.size = Pt(size)
    if bold is not None:
        run.font.bold = bold
    if italic is not None:
        run.font.italic = italic
    if color is not None:
        run.font.color.rgb = color
    # garante a fonte também para texto complexo / east-asian (evita substituição)
    rPr = run._r.get_or_add_rPr()
    for tag in ("a:ea", "a:cs"):
        el = rPr.find(qn(tag))
        if el is None:
            el = etree.SubElement(rPr, qn(tag))
        el.set("typeface", name)


def _add_rich(p, text, name, size, color, bold=False, italic=False):
    """Adiciona texto ao parágrafo; trechos entre **dois asteriscos** ficam em negrito."""
    for i, part in enumerate(re.split(r"\*\*", text)):
        if not part:
            continue
        r = p.add_run()
        r.text = part
        _set_font(r, name, size, bold=(bold or i % 2 == 1), italic=italic, color=color)


def _set_bullet(p, level, char=None):
    """Marcador real (a:buChar) com recuo pendente, como no slide de exemplo."""
    char = char or ("•" if level == 0 else "–")
    pPr = p._p.get_or_add_pPr()
    pPr.set("marL", str(int(Inches(0.35 + 0.35 * level))))
    pPr.set("indent", str(int(-Inches(0.30))))
    for tag in ("a:buNone", "a:buChar", "a:buAutoNum", "a:buFont"):
        for el in pPr.findall(qn(tag)):
            pPr.remove(el)
    bf = etree.SubElement(pPr, qn("a:buFont"))
    bf.set("typeface", "Arial")
    bc = etree.SubElement(pPr, qn("a:buChar"))
    bc.set("char", char)


def _no_bullet(p):
    pPr = p._p.get_or_add_pPr()
    pPr.set("marL", "0")
    pPr.set("indent", "0")
    for tag in ("a:buNone", "a:buChar", "a:buAutoNum", "a:buFont"):
        for el in pPr.findall(qn(tag)):
            pPr.remove(el)
    etree.SubElement(pPr, qn("a:buNone"))


def _fill_frame(tf, items, size=BODY_PT, bulleted=True, align=PP_ALIGN.LEFT, color=TEXTO,
                name=FONT_BODY, bold=False):
    tf.word_wrap = True
    for i, (text, level) in enumerate(_norm_items(items)):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = 1.0
        p.space_after = Pt(6)
        _add_rich(p, text, name, size - (2 if level >= 2 else 0), color, bold=bold)
        if bulleted or level > 0:
            _set_bullet(p, level - (0 if bulleted else 1))
        else:
            _no_bullet(p)


def _textbox(slide, x, y, w, h, items, size=BODY_PT, bulleted=True, align=PP_ALIGN.LEFT,
             color=TEXTO, anchor=MSO_ANCHOR.TOP, name=FONT_BODY, bold=False, label="texto"):
    items = _norm_items(items)
    _warn_overflow(label, items, w, h, size, bulleted)
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.vertical_anchor = anchor
    _fill_frame(tf, items, size, bulleted, align, color, name, bold)
    return tb


def _ph(slide, idx):
    for shape in slide.placeholders:
        if shape.placeholder_format.idx == idx:
            return shape
    raise KeyError(f"placeholder idx={idx} não existe neste layout")


def _remove_empty_placeholders(slide):
    for shape in list(slide.placeholders):
        if shape.has_text_frame and not shape.text_frame.text.strip():
            shape._element.getparent().remove(shape._element)


def _set_title(slide, text, name=FONT_TITLE, size=TITLE_PT, box=TITLE_BOX, color=TEXTO,
               align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.MIDDLE, bold=False):
    shape = _ph(slide, 0)
    shape.left, shape.top = Inches(box["x"]), Inches(box["y"])
    shape.width, shape.height = Inches(box["w"]), Inches(box["h"])
    tf = shape.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.text = ""
    p = tf.paragraphs[0]
    p.alignment = align
    r = p.add_run()
    r.text = text
    _set_font(r, name, size, bold=bold, color=color)
    if name == FONT_TITLE and len(text) > 38:
        warnings.warn(f"[uaippt] Título com {len(text)} caracteres pode quebrar em 2 linhas a {size} pt: '{text}'",
                      stacklevel=3)
    return shape


def _trimmed(path, pad=0.03):
    """Recorta as margens vazias (transparentes ou brancas) de PNG/JPG para o conteúdo ocupar a coluna.
    GIFs e vídeos não são alterados. Devolve caminho ou BytesIO."""
    path = Path(path)
    if path.suffix.lower() not in (".png", ".jpg", ".jpeg"):
        return str(path)
    try:
        import io
        from PIL import Image, ImageChops
        im = Image.open(path)
        rgba = im.convert("RGBA")
        alpha = rgba.split()[3]
        if alpha.getextrema()[0] < 255:                       # tem transparência
            bbox = alpha.point(lambda a: 255 if a > 8 else 0).getbbox()
        else:                                                  # fundo branco
            diff = ImageChops.difference(rgba.convert("RGB"), Image.new("RGB", im.size, (255, 255, 255)))
            bbox = diff.convert("L").point(lambda v: 255 if v > 12 else 0).getbbox()
        if not bbox:
            return str(path)
        px, py = int(im.width * pad), int(im.height * pad)
        box = (max(0, bbox[0] - px), max(0, bbox[1] - py), min(im.width, bbox[2] + px), min(im.height, bbox[3] + py))
        if box == (0, 0, im.width, im.height):
            return str(path)
        buf = io.BytesIO()
        im.crop(box).save(buf, format="PNG" if path.suffix.lower() == ".png" else "JPEG")
        buf.seek(0)
        return buf
    except Exception:
        return str(path)


def _fit_picture(slide, path, x, y, w, h, align_top=False):
    pic = slide.shapes.add_picture(_trimmed(path), Inches(x), Inches(y))
    scale = min(Inches(w) / pic.width, Inches(h) / pic.height)
    pic.width, pic.height = int(pic.width * scale), int(pic.height * scale)
    pic.left = int(Inches(x) + (Inches(w) - pic.width) / 2)
    pic.top = int(Inches(y) + (0 if align_top else (Inches(h) - pic.height) / 2))
    return pic


def _add_media(slide, path, x, y, w, h, align_top=False):
    """Imagem (PNG/JPG/GIF animado) ou vídeo (MP4…) ajustado à caixa. Devolve (left, top, w, h) em EMU."""
    path = Path(path)
    if path.suffix.lower() in (".mp4", ".mov", ".avi", ".wmv"):
        vw = min(w, h * 16 / 9)
        vh = vw * 9 / 16
        vx, vy = x + (w - vw) / 2, y + (0 if align_top else (h - vh) / 2)
        shp = slide.shapes.add_movie(str(path), Inches(vx), Inches(vy), Inches(vw), Inches(vh))
        return shp.left, shp.top, shp.width, shp.height
    pic = _fit_picture(slide, path, x, y, w, h, align_top)
    return pic.left, pic.top, pic.width, pic.height


def _caption(slide, text, right_edge_in, top_in, width_in, align=PP_ALIGN.RIGHT):
    """Crédito/legenda 12 pt logo abaixo da imagem (como 'freepik' no slide de exemplo)."""
    tb = slide.shapes.add_textbox(Inches(right_edge_in - width_in), Inches(top_in), Inches(width_in), Inches(0.3))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(0.02)
    tf.margin_top = tf.margin_bottom = Inches(0.02)
    p = tf.paragraphs[0]
    p.alignment = align
    _add_rich(p, text, FONT_BODY, CAPTION_PT, CINZA, italic=True)
    return tb


# ---------------------------------------------------------------- deck
class UaipptDeck:
    def __init__(self, template: str | Path = TEMPLATE, default_professor: str | None = None,
                 default_course: str = DEFAULT_COURSE, social_handle: str = ""):
        self.prs = Presentation(str(template))
        self.default_professor = default_professor if default_professor is not None else DEFAULT_PROFESSOR
        self.default_course = default_course
        self.social_handle = social_handle           # ex.: "@prof.ujeverson" (aparece em todos os slides)
        while len(self.prs.slides) > 0:           # template deve começar vazio
            rId = self.prs.slides._sldIdLst[0].rId
            self.prs.part.drop_rel(rId)
            del self.prs.slides._sldIdLst[0]

    # -------------------------------------------------- base
    def _new(self, layout_idx, notes=None):
        slide = self.prs.slides.add_slide(self.prs.slide_layouts[layout_idx])
        if notes:
            slide.notes_slide.notes_text_frame.text = notes
        return slide

    def _social(self, slide):
        """Ícone do Instagram + @perfil no canto inferior esquerdo (padrão do slide de exemplo)."""
        if not self.social_handle:
            return
        if SOCIAL_ICON.exists():
            slide.shapes.add_picture(str(SOCIAL_ICON), Inches(0.08), Inches(7.07), Inches(0.26), Inches(0.26))
        tb = slide.shapes.add_textbox(Inches(0.36), Inches(7.04), Inches(2.4), Inches(0.32))
        tf = tb.text_frame
        tf.margin_left = tf.margin_top = tf.margin_bottom = Inches(0.02)
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        _add_rich(tf.paragraphs[0], self.social_handle, FONT_BODY, 11, TEXTO)

    def _finish(self, slide):
        _remove_empty_placeholders(slide)
        self._social(slide)
        return slide

    # -------------------------------------------------- tipos de slide
    def cover(self, title, subtitle="", course=None, professor=None, notes=None):
        """Capa: banner, título (Poppins 36) centralizado, curso e professor(a) (Poppins 40), rodapé Fatesg."""
        s = self._new(L_CAPA, notes)
        _set_title(s, title, name=FONT_COVER, size=COVER_TITLE_PT,
                   box=dict(x=1.0, y=2.05, w=11.33, h=1.7), align=PP_ALIGN.CENTER)

        c = self.default_course if course is None else course
        p_name = self.default_professor if professor is None else professor

        # Se o nome do(a) professor(a) não foi informado e há terminal interativo, pergunta
        if p_name == "" and professor is None and sys.stdin and hasattr(sys.stdin, "isatty") and sys.stdin.isatty():
            try:
                resp = input("Nome do(a) Professor(a) para a capa (ou Enter para pular): ").strip()
                if resp:
                    p_name = resp
            except Exception:
                pass

        sub = [t for t in (subtitle, c) if t]
        sub_ph = _ph(s, 1)
        sub_ph.left, sub_ph.top, sub_ph.width, sub_ph.height = Inches(1.0), Inches(3.85), Inches(11.33), Inches(0.95)
        tf = sub_ph.text_frame
        tf.word_wrap = True
        tf.vertical_anchor = MSO_ANCHOR.TOP
        tf.text = ""
        for i, line in enumerate(sub):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.alignment = PP_ALIGN.CENTER
            _no_bullet(p)
            _add_rich(p, line, FONT_COVER, 20, CINZA)

        if p_name:
            tb = s.shapes.add_textbox(Inches(0.5), Inches(4.95), Inches(12.33), Inches(0.95))
            tf = tb.text_frame
            tf.word_wrap = True
            tf.vertical_anchor = MSO_ANCHOR.MIDDLE
            p = tf.paragraphs[0]
            p.alignment = PP_ALIGN.CENTER
            _add_rich(p, p_name.upper(), FONT_COVER, COVER_PROF_PT, TEXTO)
        return self._finish(s)

    def section(self, title, subtitle="", notes=None):
        """Divisória de seção/módulo."""
        s = self._new(L_SECAO, notes)
        shape = _ph(s, 0)
        tf = shape.text_frame
        tf.text = ""
        r = tf.paragraphs[0].add_run()
        r.text = title
        _set_font(r, FONT_TITLE, TITLE_PT, bold=True, color=AZUL)
        if subtitle:
            tf2 = _ph(s, 1).text_frame
            tf2.text = ""
            p = tf2.paragraphs[0]
            _no_bullet(p)
            _add_rich(p, subtitle, FONT_BODY, BODY_PT, CINZA)
        bar = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1.0), Inches(4.95), Inches(1.5), Inches(0.06))
        bar.fill.solid(); bar.fill.fore_color.rgb = LARANJA; bar.line.fill.background()
        return self._finish(s)

    def bullets(self, title, items, bulleted=True, notes=None):
        """Título + texto em largura total. bulleted=False → parágrafos corridos (subitens ainda levam marcador)."""
        s = self._new(L_SO_TITULO, notes)
        _set_title(s, title)
        _textbox(s, AREA["x"], AREA["y"], AREA["w"], AREA["h"], items, bulleted=bulleted, label=title)
        return self._finish(s)

    def two_columns(self, title, left, right, bulleted=True, notes=None):
        s = self._new(L_SO_TITULO, notes)
        _set_title(s, title)
        w = 5.85
        _textbox(s, AREA["x"], AREA["y"], w, AREA["h"], left, bulleted=bulleted, label=title + " (esq.)")
        _textbox(s, 6.85, AREA["y"], w, AREA["h"], right, bulleted=bulleted, label=title + " (dir.)")
        return self._finish(s)

    def comparison(self, title, left_title, left, right_title, right, notes=None):
        s = self._new(L_SO_TITULO, notes)
        _set_title(s, title)
        w = 5.85
        for x, head, items in ((AREA["x"], left_title, left), (6.85, right_title, right)):
            hb = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x + 0.1), Inches(AREA["y"]), Inches(w - 0.1), Inches(0.55))
            hb.fill.solid(); hb.fill.fore_color.rgb = AZUL; hb.line.fill.background()
            tf = hb.text_frame
            tf.vertical_anchor = MSO_ANCHOR.MIDDLE
            p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
            _add_rich(p, head, FONT_BODY, BODY_PT, BRANCO, bold=True)
            _textbox(s, x, AREA["y"] + 0.65, w, AREA["h"] - 0.65, items, label=f"{title} / {head}")
        return self._finish(s)

    def highlight(self, title, statement, items=None, notes=None):
        """Conceito/enunciado em caixa azul-clara com filete laranja + itens opcionais abaixo."""
        s = self._new(L_SO_TITULO, notes)
        _set_title(s, title)
        cpl = int((AREA["w"] - 0.8) * 72 / (BODY_PT * 0.5))
        lines = max(1, math.ceil(len(_plain(statement)) / cpl))
        box_h = min(AREA["h"] - 0.4, lines * BODY_PT * 1.25 / 72 + 0.45)
        box = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(AREA["x"] + 0.1), Inches(AREA["y"]), Inches(AREA["w"] - 0.1), Inches(box_h))
        box.fill.solid(); box.fill.fore_color.rgb = AZUL_CLARO; box.line.fill.background()
        acc = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(AREA["x"] + 0.1), Inches(AREA["y"]), Inches(0.12), Inches(box_h))
        acc.fill.solid(); acc.fill.fore_color.rgb = LARANJA; acc.line.fill.background()
        tf = box.text_frame
        tf.word_wrap = True
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        tf.margin_left, tf.margin_right = Inches(0.45), Inches(0.25)
        p = tf.paragraphs[0]; p.alignment = PP_ALIGN.LEFT
        _add_rich(p, statement, FONT_BODY, BODY_PT, AZUL, bold=True)
        if items:
            top = AREA["y"] + box_h + 0.2
            _textbox(s, AREA["x"], top, AREA["w"], AREA["y"] + AREA["h"] - top, items, label=title)
        return self._finish(s)

    def image(self, title, path, caption=None, notes=None):
        """Título + imagem grande centralizada (proporção preservada); legenda 12 pt abaixo."""
        s = self._new(L_SO_TITULO, notes)
        _set_title(s, title)
        box_h = AREA["h"] - (0.4 if caption else 0)
        l, t, w, h = _add_media(s, path, AREA["x"], AREA["y"], AREA["w"], box_h)
        if caption:
            _caption(s, caption, Emu(l + w).inches, Emu(t + h).inches + 0.03, Emu(w).inches, PP_ALIGN.CENTER)
        return self._finish(s)

    def image_text(self, title, path, text_items, caption=None, bulleted=False, media_side="right", notes=None):
        """Padrão do slide de exemplo: texto à esquerda (Calibri 24) e imagem à direita com crédito 12 pt."""
        s = self._new(L_SO_TITULO, notes)
        _set_title(s, title)
        if media_side == "left":
            img_x, text_x = AREA["x"] + 0.0, IMG_BOX["x"] - 0.9 + 0.0
            img_x = 0.66
            text_x = img_x + IMG_BOX["w"] + 0.4
        else:
            img_x, text_x = IMG_BOX["x"], AREA["x"]
        _textbox(s, text_x, AREA["y"], TEXT_W_WITH_IMG, AREA["h"], text_items, bulleted=bulleted, label=title)
        l, t, w, h = _add_media(s, path, img_x, IMG_BOX["y"], IMG_BOX["w"], IMG_BOX["h"])
        if caption:
            _caption(s, caption, Emu(l + w).inches, Emu(t + h).inches + 0.03, max(Emu(w).inches, 3.2))
        return self._finish(s)

    def two_images(self, title, paths, captions=None, text=None, notes=None):
        """Duas imagens lado a lado (slide 6 do exemplo), com uma frase de apoio acima."""
        s = self._new(L_SO_TITULO, notes)
        _set_title(s, title)
        top = AREA["y"]
        if text:
            _textbox(s, AREA["x"], top, AREA["w"], 0.9, [text] if isinstance(text, str) else text,
                     bulleted=False, label=title)
            top += 1.05
        box_w, box_h = 5.7, AREA["y"] + AREA["h"] - top - 0.4
        for k, path in enumerate(paths[:2]):
            x = 0.76 + k * (box_w + 0.5)
            l, t, w, h = _add_media(s, path, x, top, box_w, box_h)
            if captions and k < len(captions) and captions[k]:
                _caption(s, captions[k], Emu(l + w).inches, Emu(t + h).inches + 0.03, max(Emu(w).inches, 3.0))
        return self._finish(s)

    def motion_figure(self, title, path, caption=None, side_text=None, media_side="right", notes=None):
        """Figura/animação de Motion Design (Manim/Blender): PNG, GIF animado em loop ou MP4.
        Sem side_text → mídia grande centralizada; com side_text → texto à esquerda e mídia à direita."""
        path = Path(path).resolve()
        if not path.exists():
            raise FileNotFoundError(f"Arquivo de motion design não encontrado: {path}")
        if side_text:
            items = [side_text] if isinstance(side_text, str) else side_text
            return self.image_text(title, path, items, caption=caption, bulleted=True,
                                   media_side=media_side, notes=notes)
        return self.image(title, path, caption=caption, notes=notes)

    def manim_figure(self, title, path, caption=None, side_text=None, media_side="right", notes=None):
        """Atalho semântico para figuras/animações geradas via Manim."""
        return self.motion_figure(title, path, caption, side_text, media_side, notes)

    def blender_figure(self, title, path, caption=None, side_text=None, media_side="right", notes=None):
        """Atalho semântico para figuras 3D geradas via Blender."""
        return self.motion_figure(title, path, caption, side_text, media_side, notes)

    def code(self, title, code, language="python", caption=None, notes=None):
        """Bloco de código com fundo escuro e fonte monoespaçada (cores via pygments, se instalado)."""
        s = self._new(L_SO_TITULO, notes)
        _set_title(s, title)
        lines = code.expandtabs(4).rstrip("\n").split("\n")
        longest = max(len(l) for l in lines) or 1
        max_h = AREA["h"] - (0.45 if caption else 0)
        size = min(20, (max_h - 0.4) * 72 / (len(lines) * 1.2), (AREA["w"] - 0.6) * 72 / (longest * 0.6))
        size = max(10, int(size))
        box_h = min(max_h, len(lines) * size * 1.2 / 72 + 0.5)
        bg = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(AREA["x"] + 0.1), Inches(AREA["y"]), Inches(AREA["w"] - 0.1), Inches(box_h))
        bg.adjustments[0] = 0.03
        bg.fill.solid(); bg.fill.fore_color.rgb = CODE_BG; bg.line.fill.background()
        tf = bg.text_frame
        tf.word_wrap, tf.vertical_anchor = False, MSO_ANCHOR.TOP
        tf.margin_left = tf.margin_right = Inches(0.3)
        tf.margin_top = Inches(0.2)
        for i, line_tokens in enumerate(_highlight(lines, language)):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.alignment = PP_ALIGN.LEFT
            p.line_spacing = 1.0
            for text, rgb in line_tokens or [(" ", CODE_FG)]:
                r = p.add_run()
                r.text = text
                r.font.name, r.font.size, r.font.color.rgb = "Consolas", Pt(size), rgb
        acc = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(AREA["x"] + 0.1), Inches(AREA["y"] + 0.15), Inches(0.07), Inches(box_h - 0.3))
        acc.fill.solid(); acc.fill.fore_color.rgb = LARANJA; acc.line.fill.background()
        if caption:
            _caption(s, caption, AREA["x"] + AREA["w"], AREA["y"] + box_h + 0.05, AREA["w"], PP_ALIGN.CENTER)
        return self._finish(s)

    def table(self, title, header, rows, col_widths=None, notes=None):
        """Tabela Calibri com cabeçalho azul SENAI e linhas alternadas (corpo reduzido só se não couber)."""
        s = self._new(L_SO_TITULO, notes)
        _set_title(s, title)
        n_rows, n_cols = len(rows) + 1, len(header)
        row_h = min(0.7, AREA["h"] / n_rows)
        size = 20 if (row_h >= 0.6 and n_cols <= 4) else 18 if (row_h >= 0.5 and n_cols <= 5) else 16 if row_h >= 0.4 else 14
        gf = s.shapes.add_table(n_rows, n_cols, Inches(AREA["x"] + 0.1), Inches(AREA["y"]), Inches(AREA["w"] - 0.1), Inches(row_h * n_rows))
        tbl = gf.table
        if col_widths:
            total = sum(col_widths)
            for i, w in enumerate(col_widths):
                tbl.columns[i].width = Inches((AREA["w"] - 0.1) * w / total)
        for r in range(n_rows):
            values = header if r == 0 else rows[r - 1]
            for c in range(n_cols):
                cell = tbl.cell(r, c)
                cell.text = ""
                cell.fill.solid()
                cell.fill.fore_color.rgb = AZUL if r == 0 else (AZUL_CLARO if r % 2 == 0 else BRANCO)
                cell.vertical_anchor = MSO_ANCHOR.MIDDLE
                p = cell.text_frame.paragraphs[0]
                _add_rich(p, str(values[c]) if c < len(values) else "", FONT_BODY, size,
                          BRANCO if r == 0 else TEXTO, bold=(r == 0))
        return self._finish(s)

    def closing(self, title="Obrigado!", subtitle="", notes=None):
        """Encerramento (slide 2+): título Montserrat 36, subtítulo/conteúdo Calibri 24 e rodapé Fatesg."""
        s = self._new(L_CAPA, notes)
        _set_title(s, title, name=FONT_TITLE, size=TITLE_PT,
                   box=dict(x=1.0, y=2.20, w=11.33, h=1.5), align=PP_ALIGN.CENTER, bold=True, color=AZUL)
        if subtitle:
            sub_ph = _ph(s, 1)
            sub_ph.left, sub_ph.top, sub_ph.width, sub_ph.height = Inches(1.0), Inches(3.85), Inches(11.33), Inches(0.95)
            tf = sub_ph.text_frame
            tf.word_wrap = True
            tf.vertical_anchor = MSO_ANCHOR.TOP
            tf.text = ""
            p = tf.paragraphs[0]
            p.alignment = PP_ALIGN.CENTER
            _no_bullet(p)
            _add_rich(p, subtitle, FONT_BODY, BODY_PT, TEXTO)
        return self._finish(s)

    # -------------------------------------------------- saída
    def save(self, path):
        path = Path(path)
        path.parent.mkdir(parents=True, exist_ok=True)
        self.prs.save(str(path))
        return path


# Alias de retrocompatibilidade
SenaiDeck = UaipptDeck


# ---------------------------------------------------------------- realce de sintaxe (opcional)
_PALETTE = {
    "Keyword": RGBColor(0xC6, 0x78, 0xDD), "Name.Builtin": RGBColor(0x61, 0xAF, 0xEF),
    "Name.Function": RGBColor(0x61, 0xAF, 0xEF), "Name.Class": RGBColor(0xE5, 0xC0, 0x7B),
    "Literal.String": RGBColor(0x98, 0xC3, 0x79), "Literal.Number": RGBColor(0xD1, 0x9A, 0x66),
    "Comment": RGBColor(0x7F, 0x84, 0x8E), "Operator": RGBColor(0x56, 0xB6, 0xC2),
}


def _highlight(lines, language):
    try:
        from pygments import lex
        from pygments.lexers import get_lexer_by_name
        lexer = get_lexer_by_name(language)
    except Exception:
        return [[(l, CODE_FG)] if l else [] for l in lines]
    out, current = [], []
    for ttype, value in lex("\n".join(lines), lexer):
        color = CODE_FG
        for key, rgb in _PALETTE.items():
            if str(ttype).replace("Token.", "").startswith(key):
                color = rgb
                break
        parts = value.split("\n")
        for j, part in enumerate(parts):
            if j > 0:
                out.append(current)
                current = []
            if part:
                current.append((part, color))
    out.append(current)
    return out[: len(lines)]


# ---------------------------------------------------------------- CLI (roteiro JSON)
def build_from_outline(outline: dict, output: str | Path | None = None, base_dir: str | Path | None = None) -> Path:
    prof = outline.get("professor")
    if not prof and sys.stdin and hasattr(sys.stdin, "isatty") and sys.stdin.isatty():
        try:
            prof = input("Qual é o nome do(a) professor(a) responsável pela apresentação? ").strip()
        except Exception:
            prof = ""

    deck = UaipptDeck(
        template=outline.get("template", TEMPLATE),
        default_professor=prof or DEFAULT_PROFESSOR,
        default_course=outline.get("course", DEFAULT_COURSE),
        social_handle=outline.get("social_handle", ""),
    )
    for spec in outline["slides"]:
        spec = dict(spec)
        kind = spec.pop("type")
        # Resolve caminhos relativos de mídia em relação ao JSON se necessário
        if "path" in spec and base_dir:
            p = Path(spec["path"])
            if not p.is_absolute() and not p.exists():
                candidate = Path(base_dir) / p
                if candidate.exists():
                    spec["path"] = candidate
        method = getattr(deck, kind, None)
        if method is None or kind.startswith("_") or kind == "save":
            raise ValueError(f"tipo de slide desconhecido: {kind}")
        method(**spec)
    return deck.save(output or outline.get("output", "apresentacao.pptx"))


if __name__ == "__main__":
    if len(sys.argv) < 2:
        sys.exit(__doc__)
    json_path = Path(sys.argv[1]).resolve()
    data = json.loads(json_path.read_text(encoding="utf-8"))
    print(build_from_outline(data, sys.argv[2] if len(sys.argv) > 2 else None, base_dir=json_path.parent))
